from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.models.article import Article
from app.services.article_topics import TOPICS


IMPORTANCE_ORDER = {
    "high": 0,
    "medium": 1,
    "low": 2,
}

IMPORTANCE_LABELS = {
    "high": "High priority",
    "medium": "Medium priority",
    "low": "Low priority",
}


def format_published_at(
    published_at: datetime | None,
) -> str:
    if published_at is None:
        return "Unknown"

    return published_at.strftime("%Y-%m-%d %H:%M UTC")


def add_hyperlink(
    paragraph,
    text: str,
    url: str,
) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/"
        "officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")

    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)

    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def build_document(
    articles: list[Article],
    *,
    total_analyzed: int,
    total_relevant: int,
    total_not_relevant: int,
) -> Document:
    document = Document()

    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    title = document.add_heading(
        "IT / Information Security Digest",
        level=0,
    )

    title.alignment = 1

    generated_at = datetime.now(timezone.utc)

    paragraph = document.add_paragraph()
    paragraph.add_run("Generated: ").bold = True
    paragraph.add_run(
        generated_at.strftime("%Y-%m-%d %H:%M UTC")
    )

    document.add_heading("Statistics", level=1)

    statistics = [
        ("Analyzed", total_analyzed),
        ("Relevant", total_relevant),
        ("Not relevant", total_not_relevant),
    ]

    for label, value in statistics:
        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(str(value))

    sorted_articles = sorted(
        articles,
        key=lambda article: (
            TOPICS.index(article.analysis.topic)
            if article.analysis
            and article.analysis.topic in TOPICS
            else len(TOPICS),
            IMPORTANCE_ORDER.get(
                article.analysis.importance
                if article.analysis
                else "low",
                99,
            ),
            -(
                article.published_at.timestamp()
                if article.published_at
                else 0
            ),
        ),
    )

    topic_counts: dict[str, int] = {}

    for article in sorted_articles:
        if article.analysis is None:
            continue

        topic = article.analysis.topic
        topic_counts[topic] = topic_counts.get(topic, 0) + 1

    document.add_heading("Topics", level=1)

    if topic_counts:
        for topic, count in topic_counts.items():
            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(f"{topic}: ").bold = True
            paragraph.add_run(str(count))
    else:
        document.add_paragraph(
            "No relevant articles found."
        )

    if not sorted_articles:
        return document

    current_topic = None

    for index, article in enumerate(sorted_articles):
        if article.analysis is None:
            continue

        analysis = article.analysis
        topic = analysis.topic

        if topic != current_topic:
            current_topic = topic

            document.add_page_break()

            document.add_heading(
                topic,
                level=1,
            )

        document.add_heading(
            article.title,
            level=2,
        )

        source_name = (
            article.source.name
            if article.source
            else "Unknown"
        )

        paragraph = document.add_paragraph()
        paragraph.add_run("Source: ").bold = True
        paragraph.add_run(source_name)

        paragraph = document.add_paragraph()
        paragraph.add_run("Published: ").bold = True
        paragraph.add_run(
            format_published_at(article.published_at)
        )

        importance = IMPORTANCE_LABELS.get(
            analysis.importance,
            analysis.importance,
        )

        paragraph = document.add_paragraph()
        paragraph.add_run("Importance: ").bold = True
        paragraph.add_run(importance)

        paragraph = document.add_paragraph()
        paragraph.add_run("Summary:").bold = True

        document.add_paragraph(
            analysis.summary
        )

        paragraph = document.add_paragraph()
        paragraph.add_run("Why relevant:").bold = True

        document.add_paragraph(
            analysis.reason
        )

        paragraph = document.add_paragraph()
        add_hyperlink(
            paragraph,
            "Read original article",
            article.url,
        )

    return document


def save_docx(
    articles: list[Article],
    path: Path,
    *,
    total_analyzed: int,
    total_relevant: int,
    total_not_relevant: int,
) -> None:
    document = build_document(
        articles,
        total_analyzed=total_analyzed,
        total_relevant=total_relevant,
        total_not_relevant=total_not_relevant,
    )

    document.save(path)